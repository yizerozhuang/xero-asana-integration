import tkinter as tk
import tkinter.font as tkFont
from tkinter import ttk
import json
import os

from config import CONFIGURATION as conf
from utility import load_data, increment_excel_column, save, convert_float
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

SERVICE_NAME_CONVERT_MAP = {
    "Mechanical Service": "Mech",
    "CFD Service": "CFD",
    "Electrical Service": "Elec",
    "Hydraulic Service": "Hyd",
    "Fire Service": "Fire",
    "Mechanical Review": "Me",
    "Miscellaneous": "Misc",
    "Installation": "Instal",
    "Variation": "Vari"
}


def return_self_mp():
    mp_map = {
        "Quote No.": lambda data_json: data_json["Project Info"]["Project"]["Quotation Number"],
        "Proj. No.": lambda data_json: data_json["Project Info"]["Project"]["Project Number"],
        "Project detailed address": lambda data_json: data_json["Project Info"]["Project"]["Project Name"],
        "Shop Name": lambda data_json: data_json["Project Info"]["Project"]["Shop Name"],
        "Scale": lambda data_json: data_json["Project Info"]["Project"]["Proposal Type"],
        "Proj. Type": lambda data_json: data_json["Project Info"]["Project"]["Project Type"],
        "Project Status": show_status,
        "Service Been Engaged": lambda data_json: ", ".join(
            [service["Service"] for service in data_json["Project Info"]["Project"]["Service Type"].values() if
             service["Include"]]),
        "Fee Coming": lambda data_json: data_json["Email"]["Fee Coming"],
        "Fee Sent.": lambda data_json: data_json["Email"]["Fee Proposal"],
        "Fee Accept": lambda data_json: data_json["Email"]["Fee Accepted"],
        "1st": lambda data_json: data_json["Email"]["First Chase"],
        "2nd": lambda data_json: data_json["Email"]["Second Chase"],
        "3rd": lambda data_json: data_json["Email"]["Third Chase"],
        "Address To": lambda data_json: data_json["Address_to"],
        "Client Name": lambda data_json: data_json["Project Info"]["Client"]["Full Name"],
        "Client Type": lambda data_json: data_json["Project Info"]["Client"]["Contact Type"],
        "Client Company": lambda data_json: data_json["Project Info"]["Client"]["Company"],
        "Main Con Name": lambda data_json: data_json["Project Info"]["Main Contact"]["Full Name"],
        "Main Con Type": lambda data_json: data_json["Project Info"]["Main Contact"]["Contact Type"],
        "Main Con Company": lambda data_json: data_json["Project Info"]["Main Contact"]["Company"],
        "Apt/Room/Area": lambda data_json: data_json["Project Info"]["Building Features"]["Apt"],
        "Basement/Car Spots": lambda data_json: data_json["Project Info"]["Building Features"]["Basement"],
        "Feature/Notes": lambda data_json: data_json["Project Info"]["Building Features"]["Feature"],
        "INV--01": lambda data_json: show_invoice(data_json, 0),
        "INV--02": lambda data_json: show_invoice(data_json, 1),
        "INV--03": lambda data_json: show_invoice(data_json, 2),
        "INV--04": lambda data_json: show_invoice(data_json, 3),
        "INV--05": lambda data_json: show_invoice(data_json, 4),
        "INV--06": lambda data_json: show_invoice(data_json, 5),
        "Paid Fee": lambda data_json: data_json["Invoices"]["Paid Fee"],
        "Overdue Fee": lambda data_json: data_json["Invoices"]["Overdue Fee"]
    }
    invoice_lambda_function = lambda service: lambda data_json: data_json["Invoices"]["Details"][service]["in.GST"] if data_json["Invoices"]["Details"][service]["Include"] else ""




    mp_map["Total Fee"] = lambda data_json: data_json["Invoices"]["in.GST"]
    for service in conf["invoice_list"]:
        mp_map[f"{SERVICE_NAME_CONVERT_MAP[service]} Fee"] = invoice_lambda_function(service)

    bill_lambda_function = lambda service: lambda data_json: data_json["Bills"]["Details"][service]["in.GST"] if data_json["Bills"]["Details"][service]["Include"] else ""

    mp_map["Total Bill"] = lambda data_json: data_json["Bills"]["in.GST"]
    for service in conf["invoice_list"]:
        mp_map[f"{SERVICE_NAME_CONVERT_MAP[service]} Bill"] = bill_lambda_function(service)

    return mp_map

def sortby(tree, col, descending, numeric_list):
    """sort tree contents when a column header is clicked on"""
    # grab values to sort

    if col in numeric_list:
        data = [(convert_float(tree.set(child, col)), child) for child in tree.get_children('')]
    else:
        data = [(tree.set(child, col), child) for child in tree.get_children('')]
    # if the data to be sorted is numeric change to float
    # data =  change_numeric(data)
    # now sort the data in place
    data.sort(reverse=descending)
    for ix, item in enumerate(data):
        tree.move(item[1], '', ix)
    # switch the heading so it will sort in the opposite direction
    tree.heading(col, command=lambda col=col: sortby(tree, col, int(not descending), numeric_list))

def show_status(data_json):
    state = data_json["State"]
    asana_status_map = {"Pending":"06.Pending",
                        "DWG drawings":"07.DWG drawings",
                        "Done":"08.Done",
                        "Installation":"09.Installation",
                        "Construction Phase":"10.Construction"}
    if state["Asana State"] in ["Pending", "DWG drawings", "Done", "Installation", "Construction Phase"]:
        return asana_status_map[state["Asana State"]]
    elif state["Quote Unsuccessful"]:
        return "11.Quote Unsuccessful"
    elif state["Fee Accepted"]:
        return "05.Design"
    elif state["Email to Client"]:
        return "04.Chase Client"
    elif state["Generate Proposal"]:
        return "03.Email To Client"
    elif state["Set Up"]:
        return "02.Preview Fee Proposal"
    else:
        return "01.Set Up"

def show_invoice(data_json, i):
    if len(data_json["Invoices Number"][i]["Number"]) != 0:
        return data_json["Invoices Number"][i]["Number"]
    elif len(data_json["Invoices Number"][i]["Fee"]) != 0 and data_json["Invoices Number"][i]["Fee"] != "0":
        return "xxxxxx"
    else:
        return ""


class SearchBarPage(tk.Frame):
    def __init__(self, parent, app):
        tk.Frame.__init__(self, parent)
        self.parent = parent
        self.app = app

        self.main_frame = tk.Frame(self)
        self.main_frame.pack(fill=tk.BOTH, expand=1)

        self.main_canvas = tk.Canvas(self.main_frame)
        self.main_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=1, anchor="nw")
        # # self.scrollbar = ttk.Scrollbar(self.main_frame, orient=tk.VERTICAL, command=self.main_canvas.yview)
        # # self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        # # self.main_canvas.config(yscrollcommand=self.scrollbar.set)
        # self.main_canvas.bind("<Configure>", lambda e: self.main_canvas.configure(scrollregion=self.main_canvas.bbox("all")))
        # # self.app.bind("<MouseWheel>", self._on_mousewheel, add="+")
        self.main_context_frame = tk.LabelFrame(self.main_canvas, text="Main Context")
        self.main_canvas.create_window((0, 0), window=self.main_context_frame, anchor="nw")
        self.main_context_frame.bind("<Configure>", self.reset_scrollregion)
        self.conf = self.app.conf

        self.generate_convert_map()
        self.search_bar()
        self.build_tree()
        # self.generate_mp()

        self.reset()


    def generate_mp(self):
        database_dir = conf["database_dir"]
        mp_json = {}
        for dir in os.listdir(database_dir):
            # print(dir)
            if os.path.isdir(os.path.join(database_dir, dir)):
                data_dir = os.path.join(database_dir, dir, "data.json")
                data_json = json.load(open(data_dir))
                mp_json[data_json["Project Info"]["Project"]["Quotation Number"]] = {k: v(data_json) for k, v in self.mp_convert_map.items()}

        mp_dir = os.path.join(database_dir, "mp.json")

        mp_json_list = list(mp_json.keys())
        mp_json_list.sort(reverse=True)
        new_mp_json_list = {}
        for quotation in mp_json_list:
            new_mp_json_list[quotation] = mp_json[quotation]
        mp_json = new_mp_json_list
        with open(mp_dir, "w") as f:
            json_object = json.dumps(mp_json, indent=4)
            f.write(json_object)

    def reset(self):
        if len(self.app.current_quotation.get())!=0:
            save(self.app)
        self.generate_data()
        self.update_data(self.master_project)


    def generate_convert_map(self):
        self.mp_convert_map = return_self_mp()
        if not self.app.user in self.conf["admin_user_list"]:
            self.convert_map = {
                "Quote No.": lambda data_json: data_json["Project Info"]["Project"]["Quotation Number"],
                "Proj. No.": lambda data_json: data_json["Project Info"]["Project"]["Project Number"],
                "Project detailed address": lambda data_json: data_json["Project Info"]["Project"]["Project Name"],
                "Shop Name": lambda data_json: data_json["Project Info"]["Project"]["Shop Name"],
                "Scale": lambda data_json: data_json["Project Info"]["Project"]["Proposal Type"],
                "Proj. Type": lambda data_json: data_json["Project Info"]["Project"]["Project Type"],
                "Project Status": show_status,
                "Address To": lambda data_json: data_json["Address_to"],
                "Service Been Engaged": lambda data_json: ", ".join(
                    [service["Service"] for service in data_json["Project Info"]["Project"]["Service Type"].values() if
                     service["Include"]]),
                "Fee Sent.": lambda data_json: data_json["Email"]["Fee Proposal"],
                "Apt/Room/Area": lambda data_json: data_json["Project Info"]["Building Features"]["Apt"],
                "Basement/Car Spots": lambda data_json: data_json["Project Info"]["Building Features"]["Basement"],
                "Feature/Notes": lambda data_json: data_json["Project Info"]["Building Features"]["Feature"]
            }
        else:
            self.convert_map = self.mp_convert_map
        self.numeric_titles = [
            "Paid Fee",
            "Overdue Fee",
            "Total Fee ",
            "Total Bill "
        ] + [f"{SERVICE_NAME_CONVERT_MAP[service]} Fee" for service in self.conf["invoice_list"]] + [f"{SERVICE_NAME_CONVERT_MAP[service]} Bill" for service in self.conf["invoice_list"]]
        self.mp_header = list(self.convert_map.keys())


    def generate_data(self):
        database_dir = conf["database_dir"]
        res = []
        # for dir in os.listdir(database_dir):
        #     # print(dir)
        #     if os.path.isdir(os.path.join(database_dir, dir)):
        #         data_dir = os.path.join(database_dir, dir, "data.json")
        #         data_json = json.load(open(data_dir))
        #         res.append(
        #             tuple([self.convert_map[title](data_json) for title in self.mp_header])
        #         )
        mp_dir = os.path.join(database_dir, "mp.json")
        mp_json = json.load(open(mp_dir))
        for value in mp_json.values():
            res.append(tuple(value.values()))
        # res.sort(key=lambda e: e[0])
        self.master_project = res

    def search_bar(self):
        # search_bar_frame = tk.LabelFrame(self.main_context_frame)
        # search_bar_frame.pack()

        container = ttk.Frame(self.main_context_frame)
        container.pack(fill='both', expand=True)


        search_frame = tk.Frame(container)
        search_frame.grid(row=0, column=0, sticky="ew")
        self.entry = tk.Entry(search_frame, font=self.conf["font"], fg="blue", width=200)
        self.entry.grid(row=0, column=0, sticky="ew")
        tk.Button(search_frame, text="Refresh", bg="Brown", fg="white", command=self.refresh, font=self.conf["font"]).grid(row=0, column=1)
        # tk.Button(search_frame, text="Export MP Excel", bg="Brown", fg="white", command=self.export_data, font=self.conf["font"]).grid(row=0, column=2)
        tk.Label(search_frame, text="All Fee Include GST", font=self.conf["font"]+["bold"]).grid(row=0, column=2)
        self.tree = ttk.Treeview(container, height=35, columns=self.mp_header, show="headings", selectmode="browse")
        treeXScroll = ttk.Scrollbar(container, orient=tk.HORIZONTAL)
        treeXScroll.configure(command=self.main_canvas.xview)
        self.main_canvas.configure(xscrollcommand=treeXScroll.set)
        treeXScroll.grid(row=2, column=0, sticky="ew")
        # treeXScroll.pack(side=tk.BOTTOM)
        # hsb = ttk.Scrollbar(container, orient="horizontal", command=self.tree.xview)
        # self.tree.config(xscrollcommand=hsb.set)
        # hsb.config(command=self.tree.xview)
        # hsb.grid(row=2, column=0, sticky='ew')
        self.tree.grid(row=1, column=0, sticky='nsew')


        self.entry.bind("<KeyRelease>", self.check)
        self.tree.bind("<<TreeviewSelect>>", self.load_project)

    def build_tree(self):
        for col in self.mp_header:
            self.tree.heading(col, text=col.title(), command=lambda c=col: sortby(self.tree, c, 0, self.numeric_titles))
            # adjust the column's width to the header string
            self.tree.column(col, width=tkFont.Font().measure(col.title()))
    def update_data(self, data):
        self.tree.delete(*self.tree.get_children())
        for item in data:
            self.tree.insert('', 'end', values=item)
            # adjust column's width if necessary to fit each value
            # for ix, val in enumerate(item):
                # col_w = tkFont.Font().measure(val)
                # if self.tree.column(self.mp_header[ix], width=None)<col_w: self.tree.column(self.mp_header[ix], width=col_w)
    def load_project(self, event):
        self.tree.selection()
        selected = self.tree.focus()
        value = self.tree.item(selected, "values")
        # self.app.data["Project Info"]["Project"]["Quotation Number"].set(value[0])
        load_data(self.app, value[0])
    def check(self, e, search_string=None):
        if search_string is None:
            typed = self.entry.get().strip()
        else:
            typed = search_string
        if typed == "":
            data = self.master_project
        else:
            data = []
            for item in self.master_project:
                for title in item:
                    if not title is None and typed.lower() in title.lower():
                        data.append(item)
        self.update_data(set(data))
        return data


    def refresh(self):
        self.generate_data()
        self.check(None)
        sortby(self.tree, "Quote No.", True, self.numeric_titles)

    def reset_scrollregion(self, event):
        self.main_canvas.configure(scrollregion=self.main_canvas.bbox("all"))

    def export_data(self):
        report_dir = os.path.join(conf["report_dir"], "Bridge_MP.xlsx")
        cur_col = "A"
        report_wb = openpyxl.Workbook()
        report_ws = report_wb.active
        report_ws.title = "report"
        for title in self.mp_convert_map.keys():
            report_ws[f"{cur_col}1"] = title
            cur_col = increment_excel_column(cur_col)
        cur_row = 2
        for project in self.master_project:
            cur_col = "A"
            for value in project:
                report_ws[f"{cur_col}{cur_row}"] = value
                cur_col = increment_excel_column(cur_col)
            cur_row+=1
        report_wb.save(report_dir)
        # tab = Table(displayName="Bridge_MP", ref=f"A1:{cur_col}{cur_row}")
        #
        # # Add a default style with striped rows and banded columns
        # style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
        #                        showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        # tab.tableStyleInfo = style
        # report_ws.add_table(tab)


        # for project in self.master_project:
        #     new_project ={}
        #     for i, header in enumerate(self.mp_header):
        #        new_project[header] = project[i]
        #     export_json[project[0]] = new_project
        #
        # with open(os.path.join(conf["database_dir"], "mp.json"), "w") as f:
        #     json_object = json.dumps(export_json, indent=4)
        #     f.write(json_object)
    # def left_arrow(self, e):
    #     self.tree.xview_scroll(int(-500000), "units")
    # def right_arrow(self, e):
    #     self.tree.xview_scroll(int(50000), "units")

    # def reset_scrollregion(self, event):
    #     self.tree.configure(scrollregion=self.tree)